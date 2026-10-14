"""Génère Sunsoari_Codes_Douaniers.xlsx : codes SH (HS) par produit, pour la compta, la douane et les étiquettes."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

F="Arial"; NAVY="032B59"
HDR=Font(name=F,bold=True,color="FFFFFF"); FILL=PatternFill("solid",fgColor=NAVY)
B=Font(name=F); BB=Font(name=F,bold=True); T=Font(name=F,bold=True,size=14,color=NAVY)
YEL=PatternFill("solid",fgColor="FFF4CC"); SAND=PatternFill("solid",fgColor="F6EFE6")
side=Side(style="thin",color="D9CFC4"); BOX=Border(left=side,right=side,top=side,bottom=side)
W=Alignment(wrap_text=True,vertical="top")

CODES={
 "842129":("Filtre à eau pour douche (cartouche, boîtier, recharge)","Shower water filter cartridge / filter set (plastic housing, filter media)","8421.29-9029 (Sullab)"),
 "848180":("Pommeau / douchette de douche avec filtre intégré","Handheld shower head with built-in filter","8481.80-1000 (Sullab)"),
 "392690":("Housse en plastique / silicone pour douchette","Plastic/silicone protective cover for shower head","3926.90-9000 (Sullab)"),
 "842199":("Filtre de robinet et recharges (pièces de filtre)","Faucet water filter and filter refills (parts)","8421.99-9900 (Sullab) / 8421.99-9091 (Shift)"),
 "842121":("Appareil de filtration d'eau pour douche (pommeau + filtre)","Shower water filtering device (shower head with filter)","8421.21-1000 (Shift)"),
 "330730":("Capsules parfumées pour le bain et la douche (cosmétique)","Bath and shower aroma capsules (cosmetic preparation)","3307.30-2000 (Shift)"),
}
# marque, produit, sku(s), code, poids (g, Shopify), remarque
ROWS=[
 ("Onsha","Pommeau de Douche Filtrant","ONSHA-POMMEAU","848180",175,""),
 ("Onsha","Douchette Filtrante","ONSHA-DOUCHETTE","848180",155,""),
 ("Onsha","Filtres Sédiment x3 (recharges douchette)","ONSHA-SEDIMENT","842129",100,""),
 ("Onsha","Housse Personnalisable pour Douchette","ONSHA-HOUSSE-*","392690",100,"Poids à 0 sur certaines couleurs dans Shopify"),
 ("Onsha","Coque de Diffusion Home","ONSHA-COQUE-HOME","842129",None,"Poids manquant dans Shopify"),
 ("Onsha","Coque de Diffusion Nomade","ONSHA-COQUE-NOMADE","842129",None,"Poids manquant dans Shopify"),
 ("Onsha","Filtre Thermal Vitaminé (Home), 4 senteurs","ONSHA-FILTRE-*","842129",70,"Contient un gel cosmétique (fiche CPNP Sullab)"),
 ("Onsha","Filtre Thermal Vitaminé (Home) Sans Senteur","ONSHA-FILTRE-SANS-X1","842129",None,"Poids manquant dans Shopify"),
 ("Onsha","Capsule Thermale Coréenne (Nomade)","ONSHA-CAPSULE-*","842129",20,""),
 ("Onsha","Premier Rituel (Home)","ONSHA-PREMIER-*","842129",None,"Coffret : code du composant principal (filtre). Poids manquant"),
 ("Onsha","Starter Home","ONSHA-STARTER-*","842129",None,"Coffret : code du composant principal (filtres). Poids manquant"),
 ("Onsha","Set Découverte Home","ONSHA-SET-HOME","842129",445,"Code du fournisseur pour ce set"),
 ("Onsha","Set Découverte Nomade","ONSHA-SET-NOMADE","842129",220,"Poids corrigé (était saisi 220 kg au lieu de 220 g)"),
 ("Onsha","Coffret Douche Thermale Coréen (Home)","ONSHA-COFFRET-HOME","842129",710,"Coffret : code du composant principal (set de filtres)"),
 ("Onsha","Coffret Douche Thermale Coréen (Nomade)","ONSHA-COFFRET-NOMADE","842129",None,"Coffret : code du composant principal. Poids manquant"),
 ("Sullab","Filtre Robinet Lavabo","SULLAB-ROBINET","842199",None,"Poids manquant dans Shopify"),
 ("Sullab","Lot 3 Sédiments (robinet)","SULLAB-SEDIMENT-X3","842199",None,"Poids manquant dans Shopify"),
 ("SHIFT","Coffret Soin Douche Coréenne","SHIFT-COFFRET-SOIN","842121",517,"Code du fournisseur pour le Trial Kit"),
 ("SHIFT","Coffret Douche & Aromathérapie","SHIFT-COFFRET-AROMA-*","842121",1413,"Coffret : code du composant principal (pommeau filtrant)"),
 ("SHIFT","Pack 3 Filtres à Sédiments (Pure Water)","SHIFT-PUREWATER-X3","842199",78,""),
 ("SHIFT","Pack 6 Capsules Soin & Aromathérapie","SHIFT-CAPS6-*","330730",125,"Cosmétique : CPNP et étiquetage FR obligatoires dans l'UE"),
 ("SHIFT","Capsule Aromathérapie (à l'unité, brouillon)","SHIFT-*-01","330730",None,"Produit en brouillon"),
]

wb=Workbook(); ws=wb.active; ws.title="Codes par produit"; ws.sheet_view.showGridLines=False
ws["A1"]="Sunsoari · codes douaniers (SH / HS) par produit"; ws["A1"].font=T
ws["A2"]="Mis à jour le 14/10/2026. Codes enregistrés dans Shopify (fiche produit > variante > « Code SH » et « Pays d'origine »). Origine : Corée du Sud (KR) pour tous les produits."; ws["A2"].font=B
hdr=["Marque","Produit","SKU","Code SH à 6 chiffres (à mettre sur l'étiquette)","Description FR","Description EN (pour la déclaration CN22/CN23)","Code du fournisseur (Corée, 10 chiffres)","Pays d'origine","Poids (g)","Remarque"]
for i,h in enumerate(hdr,1):
    c=ws.cell(row=4,column=i,value=h); c.font=HDR; c.fill=FILL; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=BOX
ws.row_dimensions[4].height=48
for r,(br,pn,sku,code,g,rem) in enumerate(ROWS,5):
    fr,en,kr=CODES[code]
    vals=[br,pn,sku,code,fr,en,kr,"Corée du Sud (KR)",g if g else "à peser",rem]
    for i,v in enumerate(vals,1):
        c=ws.cell(row=r,column=i,value=v); c.font=BB if i==4 else B; c.alignment=W; c.border=BOX
        if i==4: c.number_format="@"; c.fill=SAND
        if i==9 and not g: c.fill=YEL
    ws.row_dimensions[r].height=32
for col,w in zip("ABCDEFGHIJ",[9,40,24,16,38,44,26,16,10,44]): ws.column_dimensions[col].width=w
ws.freeze_panes="C5"

ws2=wb.create_sheet("Lexique"); ws2.sheet_view.showGridLines=False; ws2.column_dimensions["A"].width=30; ws2.column_dimensions["B"].width=100
lex=[("C'est quoi ce numéro ?",""),
 ("Code SH = code HS","Le même code, en français (Système Harmonisé) ou en anglais (Harmonized System). 6 chiffres, identiques dans tous les pays du monde : France, UE, Suisse, Afrique, Corée. C'est celui à mettre sur une étiquette d'expédition."),
 ("Code NC (8 chiffres)","Version européenne : les 6 chiffres SH + 2 chiffres ajoutés par l'UE. Utile pour un dédouanement d'import en France. À confirmer par ton transitaire."),
 ("Code TARIC (10 chiffres)","Version européenne la plus détaillée, pour calculer les droits à l'import dans l'UE."),
 ("Code du fournisseur (10 chiffres)","Le code coréen écrit sur les factures Sullab et Shift (ex. 8421.29-9029). Seuls les 6 premiers chiffres sont valables partout."),
 ("Code EAN","Ce n'est PAS un code douanier : c'est le code-barres à 13 chiffres (ex. 8809532133116 chez Shift). Il sert à Shopify, Google et aux magasins."),
 ("",""),
 ("Envoyer hors UE (Suisse, Afrique…)",""),
 ("Étiquette / CN22-CN23","Pour chaque article : description EN, quantité, valeur, poids, code SH à 6 chiffres, pays d'origine KR. Envoi de marchandise vendue : type « Vente de marchandises » (Sale of goods)."),
 ("Suisse","La cliente peut payer la TVA suisse (8,1 %) et des frais de dédouanement à la livraison si le transporteur les lui facture. Préviens-la sur ta page livraison."),
 ("Afrique","Les droits et taxes varient selon le pays. Le code SH à 6 chiffres reste le même ; le reste du code dépend du pays de destination."),
 ("Cosmétiques (code 3307)","Les capsules Shift sont des cosmétiques : certains pays peuvent demander une fiche de composition (INCI) ou une fiche de sécurité. Garde-les à portée de main."),
 ("Coffrets","Un coffret prend le code de son composant principal (règle douanière 3b). Les codes de coffrets sont donc à confirmer par ton transitaire."),
]
for r,(a,b) in enumerate(lex,1):
    ws2.cell(row=r,column=1,value=a).font=BB if (b=="" and a) else BB
    c=ws2.cell(row=r,column=2,value=b); c.font=B; c.alignment=W
    if b=="" and a: ws2.cell(row=r,column=1).font=T
    ws2.row_dimensions[r].height=30 if len(b)<110 else 45
wb.save(sys.argv[1]); print("ok")
