"""Génère Sunsoari_Marges.xlsx : tarifs fournisseurs Onsha/Sullab + Shift, coûts rendus, marges."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="8C6A4F")
YEL = PatternFill("solid", fgColor="FFFF00")
SOFT = PatternFill("solid", fgColor="F6EFE6")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D9CFC4")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00 €;(#,##0.00 €);"-"'
USD = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
COEF = '0.00"x";(0.00"x");"-"'

wb = Workbook()

def sheet(title):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    return ws

def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.alignment, c.border = HDR, HDR_FILL, Alignment(wrap_text=True, vertical="center", horizontal="center"), BOX
    ws.row_dimensions[row].height = 45
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

def put(ws, ref, value, font=BLACK, fmt=None, fill=None, note=None):
    c = ws[ref]
    c.value, c.font = value, font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if note: c.comment = Comment(note, "Claude")
    c.border = BOX
    return c

# ------------------------------------------------------------------ Mode d'emploi
ws = wb.active
ws.title = "Mode d'emploi"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110
lines = [
    ("Sunsoari · Tableau de marges et de gestion", TITLE),
    ("Version du 11/10/2026 · prix de vente repris de Shopify le 11/10/2026 · tarifs Onsha/Sullab : fiche FOB USD Sullab Co., Ltd.", BLACK),
    ("", BLACK),
    ("CODE COULEUR", BOLD),
    ("Texte bleu = chiffre que tu peux modifier.  Fond jaune = à remplir ou à vérifier en priorité.  Texte noir = calcul automatique (ne pas toucher).  Texte vert = repris d'un autre onglet.", BLACK),
    ("", BLACK),
    ("DANS QUEL ORDRE REMPLIR", BOLD),
    ("1. Paramètres : ton statut TVA, le taux du dollar, tes frais Shopify, ta livraison, ton emballage, tes cotisations.", BLACK),
    ("2. Import : les vrais frais de ta dernière commande fournisseur (fret, transitaire, banque). Ça calcule ton « coefficient d'approche » = combien te coûte vraiment 1 € de marchandise une fois arrivée chez toi.", BLACK),
    ("3. Tarifs Shift : à compléter avec ta fiche Shift (cases jaunes).", BLACK),
    ("4. Marges produits : prix de vente, coût de revient, marge par produit et marge si la cliente n'achète que ce produit.", BLACK),
    ("5. Simulateur commande : compose un panier et vois ce qu'il te reste vraiment.", BLACK),
    ("6. Seuil de rentabilité : combien de commandes par mois pour couvrir tes frais fixes.", BLACK),
    ("", BLACK),
    ("À SAVOIR", BOLD),
    ("• FOB = le prix du fournisseur s'arrête au port coréen. Le transport international, l'assurance, le dédouanement, les droits et la TVA à l'import sont à ta charge : c'est pour ça que le coût réel est plus élevé que le prix de la fiche.", BLACK),
    ("• Droits de douane : l'accord de libre-échange UE-Corée permet 0 % de droits si le fournisseur fournit une déclaration d'origine sur la facture. Sans elle, compte en général entre 1,7 % et 6,5 % selon le code douanier. À faire confirmer par ton transitaire ou ton transporteur.", BLACK),
    ("• TVA : si tu es en franchise de TVA (micro-entreprise sous le seuil), tu ne factures pas de TVA mais tu ne récupères pas non plus la TVA payée à l'import : elle devient un coût. Si tu factures la TVA, tu la récupères, mais ton chiffre d'affaires réel est le prix TTC divisé par 1,2.", BLACK),
    ("• La correspondance entre produits Shopify et références fournisseur est marquée « à confirmer » quand je l'ai déduite des photos ou des descriptions. Vérifie-la une fois avec tes factures.", BLACK),
    ("• Les taux (frais Shopify, cotisations URSSAF, taux de change) changent : vérifie-les dans ton admin Shopify et sur autoentrepreneur.urssaf.fr. Ce tableau ne remplace pas l'avis de ton expert-comptable.", BLACK),
]
for i, (t, f) in enumerate(lines, 2):
    c = ws.cell(row=i, column=2, value=t)
    c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")

# ------------------------------------------------------------------ Paramètres
P = sheet("Paramètres")
for col, w in zip("ABCD", [3, 48, 16, 80]):
    P.column_dimensions[col].width = w
put(P, "B2", "Paramètres généraux", TITLE).border = Border()
header_row = 4
for i, lab in enumerate(["", "Paramètre", "Valeur", "Explication / source"], 1):
    if i > 1:
        c = P.cell(row=header_row, column=i, value=lab)
        c.font, c.fill, c.border = HDR, HDR_FILL, BOX
params = [
    # key, label, value, fmt, note, yellow
    ("statut", "Statut TVA", "Franchise", None, "« Franchise » si tu ne factures pas la TVA (micro-entreprise sous le seuil), « Assujettie » si tu factures la TVA. À confirmer avec ta situation réelle.", True),
    ("tva", "Taux de TVA", 0.20, PCT, "Taux normal en France.", False),
    ("usd", "Taux de change : 1 USD = … €", 0.86, '0.0000 €', "Indicatif. Mets le taux réellement appliqué par ta banque le jour du paiement (relevé bancaire).", True),
    ("krw", "Taux de change : 1 KRW = … €", 0.00063, '0.000000 €', "Indicatif, utile si Shift facture en wons. À mettre à jour.", False),
    ("palier", "Palier de prix Onsha/Sullab utilisé", "Carton", None, "« Carton » = prix au carton (petites quantités). « MOQ » = prix à partir d'environ 1 000 pièces par référence.", False),
    ("douane", "Droits de douane par défaut", 0.0, PCT, "0 % avec l'accord UE-Corée si une déclaration d'origine figure sur la facture. Sinon souvent entre 1,7 % et 6,5 % selon le code douanier (à confirmer par ton transitaire).", True),
    ("paypct", "Frais de paiement Shopify (%)", 0.015, PCT, "Exemple : Shopify Payments, forfait Basic, carte européenne. Vérifie dans Shopify > Paramètres > Paiements (PayPal, Klarna et cartes hors UE coûtent plus cher).", True),
    ("payfix", "Frais de paiement Shopify (fixe par commande)", 0.25, EUR, "Même source que ci-dessus.", True),
    ("cot", "Cotisations sociales (% du CA encaissé)", 0.123, PCT, "Micro-entreprise, vente de marchandises : taux indicatif à vérifier sur autoentrepreneur.urssaf.fr (+1 % si versement libératoire de l'impôt). Mets 0 % si tu es en société.", True),
    ("livr", "Coût moyen d'expédition d'un colis", 5.50, EUR, "Ce que TU paies au transporteur (Colissimo, Mondial Relay…). Exemple à remplacer par ton tarif réel.", True),
    ("livrfact", "Livraison payée par la cliente (moyenne TTC)", 0.0, EUR, "0 si livraison offerte. Sinon le montant moyen facturé.", False),
    ("emb", "Emballage par colis", 1.00, EUR, "Carton, papier de soie, carte, scotch, étiquette. Exemple à remplacer.", True),
    ("pub", "Coût publicitaire moyen par commande", 0.0, EUR, "0 tant que tu ne fais pas de pub. Ensuite : budget pub du mois ÷ nombre de commandes du mois.", False),
    ("cible", "Coefficient multiplicateur visé", 3.0, COEF, "Prix de vente TTC ÷ coût de revient. Repère courant en beauté : entre 2,5 et 3,5. Sert à calculer un prix plancher.", False),
    ("margemin", "Marge minimum acceptable (% du CA)", 0.40, PCT, "En dessous, le produit est signalé « marge faible ».", False),
]
PR = {}
r = 5
for key, lab, val, fmt, note, yellow in params:
    put(P, f"B{r}", lab)
    put(P, f"C{r}", val, BLUE, fmt, YEL if yellow else None)
    c = put(P, f"D{r}", note); c.alignment = WRAP
    PR[key] = f"Paramètres!$C${r}"
    P.row_dimensions[r].height = 32
    r += 1
dv = DataValidation(type="list", formula1='"Franchise,Assujettie"', allow_blank=False); P.add_data_validation(dv); dv.add(PR["statut"].split("!")[1].replace("$", ""))
dv2 = DataValidation(type="list", formula1='"Carton,MOQ"', allow_blank=False); P.add_data_validation(dv2); dv2.add(PR["palier"].split("!")[1].replace("$", ""))
TVA_RECUP = f'({PR["statut"]}="Assujettie")'

# ------------------------------------------------------------------ Import
I = sheet("Import")
for col, w in zip("ABCDE", [3, 46, 18, 18, 70]):
    I.column_dimensions[col].width = w
put(I, "B2", "Coût réel d'une commande fournisseur (coefficient d'approche)", TITLE).border = Border()
c = I["B3"]; c.value = "Remplace les valeurs d'exemple (jaunes) par celles de ta dernière facture fournisseur et de ta facture transporteur/transitaire."; c.font = BLACK
for i, lab in enumerate(["Ligne", "Onsha / Sullab", "Shift", "Explication"], 2):
    c = I.cell(row=5, column=i, value=lab); c.font, c.fill, c.border = HDR, HDR_FILL, BOX
rows = [
    ("Valeur de la marchandise (devise du fournisseur)", 2000, 1000, NUM, True, "Total de la facture fournisseur (pro forma). Exemple."),
    ("Devise de la facture", "USD", "USD", None, True, "USD, EUR ou KRW."),
    ("Taux de change appliqué", "rate", "rate", '0.000000', False, "Repris des Paramètres."),
    ("Valeur de la marchandise en €", "val", "val", EUR, False, ""),
    ("Fret international (€)", 400, 250, EUR, True, "Transport Corée → France (avion ou mer). Exemple à remplacer par le vrai devis."),
    ("Assurance transport (€)", 0, 0, EUR, False, "Souvent incluse ou optionnelle."),
    ("Frais de dédouanement / transitaire (€)", 90, 60, EUR, True, "Frais de dossier du transporteur ou du transitaire. Exemple."),
    ("Frais bancaires du virement international (€)", 30, 30, EUR, True, "Paiement T/T en 2 fois chez Sullab (50 % avant, 50 % après expédition) : compte 2 virements."),
    ("Taux de droits de douane", "duty", "duty", PCT, False, "Repris des Paramètres (modifiable ici si un fournisseur a un taux différent)."),
    ("Droits de douane (€)", "dutyeur", "dutyeur", EUR, False, "Calculés sur marchandise + fret + assurance."),
    ("TVA à l'import (€)", "tvaimp", "tvaimp", EUR, False, "20 % sur (marchandise + fret + assurance + droits). Récupérable seulement si tu es assujettie."),
    ("Total payé pour cette commande (€)", "total", "total", EUR, False, ""),
    ("Coefficient d'approche hors TVA", "cht", "cht", '0.000', False, "Coût arrivé ÷ valeur de la marchandise, sans la TVA import."),
    ("Coefficient d'approche TVA import comprise", "cttc", "cttc", '0.000', False, "Idem en comptant la TVA import comme un coût."),
    ("COEFFICIENT RETENU", "cret", "cret", '0.000', False, "Selon ton statut TVA. Il est appliqué à tous les prix de cet onglet fournisseur."),
]
for i, (lab, a, b, fmt, yellow, note) in enumerate(rows):
    rr = 6 + i
    put(I, f"B{rr}", lab, BOLD if lab.startswith("COEFF") else BLACK)
    for col, v in (("C", a), ("D", b)):
        ref = f"{col}{rr}"
        formulas = {
            "rate": f'=IF({col}7="USD",{PR["usd"]},IF({col}7="KRW",{PR["krw"]},1))',
            "val": f"={col}6*{col}8",
            "duty": f"={PR['douane']}",
            "dutyeur": f"=({col}9+{col}10+{col}11)*{col}14",
            "tvaimp": f"=({col}9+{col}10+{col}11+{col}15)*{PR['tva']}",
            "total": f"=SUM({col}9:{col}13)+{col}15+{col}16",
            "cht": f"=IF({col}9=0,1,({col}17-{col}16)/{col}9)",
            "cttc": f"=IF({col}9=0,1,{col}17/{col}9)",
            "cret": f"=IF({TVA_RECUP},{col}18,{col}19)",
        }
        if v in formulas:
            put(I, ref, formulas[v], GREEN if v in ("rate", "duty") else (BOLD if v == "cret" else BLACK), fmt)
        else:
            put(I, ref, v, BLUE, fmt, YEL if yellow else None)
    c = put(I, f"E{rr}", note); c.alignment = WRAP
dv3 = DataValidation(type="list", formula1='"USD,EUR,KRW"'); I.add_data_validation(dv3); dv3.add("C7:D7")
COEF_ON, COEF_SH = "Import!$C$20", "Import!$D$20"

# ------------------------------------------------------------------ Tarifs Onsha-Sullab
T = sheet("Tarifs Onsha-Sullab")
put(T, "A1", "Tarifs Sullab Co., Ltd. (Onsha / Sullab) · FOB Corée, en USD", TITLE).border = Border()
T["A2"] = "Source : fiche « Onsha FOB Price USD » (Sullab Co., Ltd.). Paiement 50 % T/T à la commande, solde après expédition. Carton 55 × 34 × 19 cm, 60 cartons max par palette."
T["A2"].font = BLACK
cols = ["Réf.", "Produit (fiche fournisseur)", "Correspondance Sunsoari", "Prix unitaire au carton (USD)", "Prix d'un carton (USD)", "Pièces par carton",
        "Prix unitaire MOQ (USD)", "Quantité MOQ (pcs)", "Total MOQ (USD)", "Prix public Corée (USD)", "Poids carton (kg)", "Remarques",
        "Prix retenu (USD)", "Coût d'achat (€)", "Coût de revient rendu (€)"]
header(T, 4, cols, [6, 26, 34, 13, 12, 10, 13, 11, 13, 12, 10, 30, 12, 12, 14])
sullab = [
    (1, "Shower Head", "Douchette filtrante Onsha (à confirmer)", 6.50, 40, 5.50, 1000, 15.50, 7.5, ""),
    (2, "Shower Head Cover", "Housse personnalisable pour douchette", 3.70, 50, 3.20, 1000, 6.54, 5.5, ""),
    (3, "Shower Head Filter", "Filtres sédiment x3 (recharges douchette)", 4.20, 54, 3.80, 1026, 9.50, 6.5, "1 paquet = 3 filtres"),
    (4, "PVC Shower Tube 1.5 M", "Non vendu actuellement (flexible 1,5 m)", 6.00, 24, 5.00, 1008, 15.00, 8.0, "Blanc, noir, gris, bleu, violet, rose"),
    (5, "PVC Shower Tube 2.0 M", "Non vendu actuellement (flexible 2 m)", 6.50, 24, 5.20, 1008, 16.00, 9.0, "Blanc, noir, gris, bleu, violet, rose"),
    (6, "Hot-Spring Filter Set (Mini)", "Set Découverte Nomade (coque + 4 capsules)", 10.20, 18, 9.20, 1008, 29.00, 5.0, ""),
    (7, "Hot-Spring Refill Filter: Wood", "Capsule thermale Nomade · Hinoki", 2.30, 200, 1.80, 1000, 6.54, 5.5, ""),
    (8, "Hot-Spring Refill Filter: Ocean", "Capsule thermale Nomade · Océan", 2.30, 200, 1.80, 1000, 6.54, 5.5, ""),
    (9, "Hot-Spring Refill Filter: Pine Tree", "Capsule thermale Nomade · Forêt de Pin", 2.30, 200, 1.80, 1000, 6.54, 5.5, ""),
    (10, "Hot-Spring Refill Filter: Flower", "Capsule thermale Nomade · Fleur de Prunier", 2.30, 200, 1.80, 1000, 6.54, 5.5, ""),
    (11, "Hot-Spring Filter case (Mini)", "Coque de diffusion Nomade", 2.30, 200, 1.80, 1000, 6.15, 10.5, ""),
    (12, "Compact Shower Head", "Pommeau de douche filtrant Onsha (à confirmer)", 6.20, 42, 5.20, 1008, 12.50, 8.5, ""),
    (13, "Hot-Spring Filter Set (Home)", "Set Découverte Home (coque + 4 filtres)", 20.00, 12, 18.50, 1008, 45.00, 6.5, ""),
    (14, "Hot-Spring Refill Filter: Wood (Home)", "Filtre thermal Home · Hinoki", 4.70, 72, 4.20, 1008, 11.00, 6.5, ""),
    (15, "Hot-Spring Refill Filter: Ocean (Home)", "Filtre thermal Home · Océan", 4.70, 72, 4.20, 1008, 11.00, 6.5, ""),
    (16, "Hot-Spring Refill Filter: Pine Tree (Home)", "Filtre thermal Home · Forêt de Pin", 4.70, 72, 4.20, 1008, 11.00, 6.5, ""),
    (17, "Hot-Spring Refill Filter: Flower (Home)", "Filtre thermal Home · Fleur de Prunier", 4.70, 72, 4.20, 1008, 11.00, 6.5, ""),
    (18, "Hot-Spring Filter case (HOME)", "Coque de diffusion Home (à confirmer)", 2.30, 72, 1.80, 1008, 9.00, 8.5, ""),
    (19, "Faucet Filter with case (HOME)", "Sullab · Filtre robinet lavabo", 5.40, 72, 4.60, 1008, 15.00, 9.0, "6 adaptateurs et outil inclus"),
    (20, "Hot-Spring Filter case (HOME) · paquet de 3", "Sullab · Lot 3 sédiments robinet (à confirmer)", 2.30, 212, 1.90, 1060, 9.00, 8.5, "1 paquet = 3 filtres"),
]
ON = {}
for i, (ref, name, corr, pu, qty, pmoq, qmoq, retail, kg, rem) in enumerate(sullab):
    r = 5 + i
    ON[ref] = r
    put(T, f"A{r}", ref)
    put(T, f"B{r}", name)
    put(T, f"C{r}", corr, BLUE, fill=YEL if "confirmer" in corr else None)
    put(T, f"D{r}", pu, BLUE, USD)
    put(T, f"E{r}", f"=D{r}*F{r}", BLACK, USD)
    put(T, f"F{r}", qty, BLUE, NUM)
    put(T, f"G{r}", pmoq, BLUE, USD)
    put(T, f"H{r}", qmoq, BLUE, NUM)
    put(T, f"I{r}", f"=G{r}*H{r}", BLACK, USD)
    put(T, f"J{r}", retail, BLUE, USD)
    put(T, f"K{r}", kg, BLUE, '0.0')
    put(T, f"L{r}", rem)
    put(T, f"M{r}", f'=IF({PR["palier"]}="MOQ",G{r},D{r})', BLACK, USD)
    put(T, f"N{r}", f'=M{r}*{PR["usd"]}', BLACK, EUR)
    put(T, f"O{r}", f"=N{r}*{COEF_ON}", BOLD, EUR)
T.freeze_panes = "D5"

# ------------------------------------------------------------------ Tarifs Shift
S = sheet("Tarifs Shift")
put(S, "A1", "Tarifs SHIFT · à compléter avec ta fiche fournisseur", TITLE).border = Border()
S["A2"] = "Remplis les cases jaunes (prix unitaire et devise). Ajoute des lignes si ta fiche contient d'autres références : copie une ligne existante pour garder les formules."
S["A2"].font = BLACK
header(S, 4, ["Réf.", "Composant", "Utilisé dans", "Prix unitaire fournisseur", "Devise", "Quantité minimum", "Remarques", "Taux de change", "Coût d'achat (€)", "Coût de revient rendu (€)"],
       [6, 34, 40, 14, 9, 11, 30, 12, 13, 14])
shift = [
    ("S1", "Pommeau de douche filtrant SHIFT", "Coffret Soin, Coffret Aromathérapie"),
    ("S2", "Flexible silicone 2 m", "Coffret Aromathérapie"),
    ("S3", "Filtre à sédiments Pure Water (1 pièce)", "Pack 3 filtres, coffrets"),
    ("S4", "Capsule soin & aromathérapie (1 pièce)", "Pack 6 capsules, coffrets"),
    ("S5", "Autre référence (à renommer)", ""),
    ("S6", "Autre référence (à renommer)", ""),
]
SH = {}
for i, (ref, name, used) in enumerate(shift):
    r = 5 + i
    SH[ref] = r
    put(S, f"A{r}", ref)
    put(S, f"B{r}", name, BLUE)
    put(S, f"C{r}", used)
    put(S, f"D{r}", None, BLUE, '#,##0.00', YEL)
    put(S, f"E{r}", "USD", BLUE, fill=YEL)
    put(S, f"F{r}", None, BLUE, NUM)
    put(S, f"G{r}", "")
    put(S, f"H{r}", f'=IF(E{r}="USD",{PR["usd"]},IF(E{r}="KRW",{PR["krw"]},1))', GREEN, '0.000000')
    put(S, f"I{r}", f"=D{r}*H{r}", BLACK, EUR)
    put(S, f"J{r}", f"=I{r}*{COEF_SH}", BOLD, EUR)
dv4 = DataValidation(type="list", formula1='"USD,EUR,KRW"'); S.add_data_validation(dv4); dv4.add("E5:E10")

# ------------------------------------------------------------------ Marges produits
M = sheet("Marges produits")
put(M, "A1", "Marge par produit Sunsoari", TITLE).border = Border()
M["A2"] = "Prix de vente repris de Shopify le 11/10/2026 (modifiables en bleu). La colonne « Composition » explique quelles références fournisseur forment chaque produit."
M["A2"].font = BLACK
mcols = ["Marque", "Produit Shopify", "SKU", "Composition (réf. fournisseur)", "Prix de vente TTC (€)", "Coût d'achat fournisseur (€)",
         "Coût de revient rendu (€)", "CA encaissé (€)", "Frais de paiement % (€)", "Cotisations (€)", "Marge produit (€)",
         "Taux de marge (% du CA)", "Coefficient (prix TTC ÷ coût)", "Marge si commande de cet article seul (€)", "Prix plancher au coefficient visé (€)", "Alerte"]
header(M, 4, mcols, [10, 40, 22, 44, 12, 13, 13, 12, 12, 12, 12, 12, 13, 15, 14, 26])
ONS = "'Tarifs Onsha-Sullab'"
SHS = "'Tarifs Shift'"
def on(parts):  # parts: list of (qty, ref)
    n = "+".join(f"{q}*{ONS}!$N${ON[r]}" if q != 1 else f"{ONS}!$N${ON[r]}" for q, r in parts)
    g = "+".join(f"{q}*{ONS}!$O${ON[r]}" if q != 1 else f"{ONS}!$O${ON[r]}" for q, r in parts)
    return "=" + n, "=" + g
def sh(parts):
    n = "+".join(f"{q}*{SHS}!$I${SH[r]}" if q != 1 else f"{SHS}!$I${SH[r]}" for q, r in parts)
    g = "+".join(f"{q}*{SHS}!$J${SH[r]}" if q != 1 else f"{SHS}!$J${SH[r]}" for q, r in parts)
    return "=" + n, "=" + g
products = [
    ("Onsha", "Pommeau de Douche Filtrant", "ONSHA-POMMEAU", 23.90, "Réf. 12 (à confirmer)", on([(1, 12)])),
    ("Onsha", "Douchette Filtrante", "ONSHA-DOUCHETTE", 23.90, "Réf. 1 (à confirmer)", on([(1, 1)])),
    ("Onsha", "Filtres Sédiment x3 (recharges douchette)", "ONSHA-SEDIMENT", 13.90, "Réf. 3 (paquet de 3)", on([(1, 3)])),
    ("Onsha", "Housse Personnalisable pour Douchette", "ONSHA-HOUSSE-*", 9.90, "Réf. 2", on([(1, 2)])),
    ("Onsha", "Coque de Diffusion Home", "ONSHA-COQUE-HOME", 19.90, "Réf. 18 (à confirmer)", on([(1, 18)])),
    ("Onsha", "Coque de Diffusion Nomade", "ONSHA-COQUE-NOMADE", 9.90, "Réf. 11", on([(1, 11)])),
    ("Onsha", "Filtre Thermal Vitaminé (Home)", "ONSHA-FILTRE-*", 12.90, "Réf. 14 à 17", on([(1, 14)])),
    ("Onsha", "Filtre Thermal Vitaminé (Home) Sans Senteur", "ONSHA-FILTRE-SANS-X1", 12.90, "Absent de la fiche : même prix que réf. 14 supposé", on([(1, 14)])),
    ("Onsha", "Capsule Thermale Coréenne (Nomade)", "ONSHA-CAPSULE-*", 6.90, "Réf. 7 à 10", on([(1, 7)])),
    ("Onsha", "Premier Rituel (Home)", "ONSHA-PREMIER-*", 29.90, "Coque Home (18) + 1 filtre Home (14)", on([(1, 18), (1, 14)])),
    ("Onsha", "Starter Home", "ONSHA-STARTER-*", 49.90, "Pommeau (12) + coque Home (18) + 1 filtre Home (14)", on([(1, 12), (1, 18), (1, 14)])),
    ("Onsha", "Set Découverte Home", "ONSHA-SET-HOME", 59.90, "Réf. 13 (set coque + 4 filtres)", on([(1, 13)])),
    ("Onsha", "Set Découverte Nomade", "ONSHA-SET-NOMADE", 29.90, "Réf. 6 (set coque + 4 capsules)", on([(1, 6)])),
    ("Onsha", "Coffret Douche Thermale Coréen (Home)", "ONSHA-COFFRET-HOME", 79.90, "Pommeau (12) + set Home (13)", on([(1, 12), (1, 13)])),
    ("Onsha", "Coffret Douche Thermale Coréen (Nomade)", "ONSHA-COFFRET-NOMADE", 59.90, "Douchette (1) + sédiments x3 (3) + set Nomade (6)", on([(1, 1), (1, 3), (1, 6)])),
    ("Sullab", "Filtre Robinet Lavabo", "SULLAB-ROBINET", 25.90, "Réf. 19", on([(1, 19)])),
    ("Sullab", "Lot 3 Sédiments (robinet)", "SULLAB-SEDIMENT-X3", 11.90, "Réf. 20 (à confirmer)", on([(1, 20)])),
    ("SHIFT", "Coffret Soin Douche Coréenne", "SHIFT-COFFRET-SOIN", 59.90, "Pommeau (S1) + 1 filtre (S3) + 1 capsule (S4)", sh([(1, "S1"), (1, "S3"), (1, "S4")])),
    ("SHIFT", "Coffret Douche & Aromathérapie", "SHIFT-COFFRET-AROMA-*", 109.90, "Pommeau (S1) + flexible (S2) + 3 filtres (S3) + 6 capsules (S4)", sh([(1, "S1"), (1, "S2"), (3, "S3"), (6, "S4")])),
    ("SHIFT", "Pack 3 Filtres à Sédiments (Pure Water)", "SHIFT-PUREWATER-X3", 25.90, "3 filtres (S3)", sh([(3, "S3")])),
    ("SHIFT", "Pack 6 Capsules Soin & Aromathérapie", "SHIFT-CAPS6-*", 32.90, "6 capsules (S4)", sh([(6, "S4")])),
]
first = 5
for i, (brand, name, sku, price, comp, (fn, gn)) in enumerate(products):
    r = first + i
    put(M, f"A{r}", brand)
    put(M, f"B{r}", name)
    put(M, f"C{r}", sku)
    c = put(M, f"D{r}", comp, fill=YEL if ("confirmer" in comp or "supposé" in comp) else None); c.alignment = WRAP
    put(M, f"E{r}", price, BLUE, EUR)
    put(M, f"F{r}", fn, GREEN, EUR)
    put(M, f"G{r}", gn, GREEN, EUR)
    put(M, f"H{r}", f"=IF({TVA_RECUP},E{r}/(1+{PR['tva']}),E{r})", BLACK, EUR)
    put(M, f"I{r}", f"=E{r}*{PR['paypct']}", BLACK, EUR)
    put(M, f"J{r}", f"=H{r}*{PR['cot']}", BLACK, EUR)
    put(M, f"K{r}", f"=IF(G{r}=0,0,H{r}-G{r}-I{r}-J{r})", BOLD, EUR)
    put(M, f"L{r}", f"=IF(H{r}=0,0,K{r}/H{r})", BLACK, PCT)
    put(M, f"M{r}", f"=IF(G{r}=0,0,E{r}/G{r})", BLACK, COEF)
    put(M, f"N{r}", f"=IF(G{r}=0,0,K{r}-{PR['payfix']}-{PR['livr']}-{PR['emb']}-{PR['pub']}+IF({TVA_RECUP},{PR['livrfact']}/(1+{PR['tva']}),{PR['livrfact']})*(1-{PR['cot']}))", BOLD, EUR)
    put(M, f"O{r}", f"=G{r}*{PR['cible']}", BLACK, EUR)
    put(M, f"P{r}", f'=IF(G{r}=0,"Prix fournisseur manquant",IF(N{r}<0,"Perte si vendu seul",IF(L{r}<{PR["margemin"]},"Marge faible","OK")))', BLACK)
    M.row_dimensions[r].height = 30
last = first + len(products) - 1
M.freeze_panes = "C5"
from openpyxl.formatting.rule import FormulaRule
M.conditional_formatting.add(f"P{first}:P{last}", FormulaRule(formula=[f'P{first}="OK"'], fill=PatternFill("solid", fgColor="D8EBD3")))
M.conditional_formatting.add(f"P{first}:P{last}", FormulaRule(formula=[f'P{first}<>"OK"'], fill=PatternFill("solid", fgColor="F8D7D2")))

# ------------------------------------------------------------------ Simulateur commande
C = sheet("Simulateur commande")
for col, w in zip("ABCDEF", [3, 44, 10, 14, 14, 50]):
    C.column_dimensions[col].width = w
put(C, "B2", "Simulateur de commande", TITLE).border = Border()
C["B3"] = "Choisis jusqu'à 6 produits dans la liste et leur quantité : tu vois ce qu'il te reste après tous les frais."
C["B3"].font = BLACK
for i, lab in enumerate(["Produit", "Quantité", "Prix TTC total (€)", "Coût de revient total (€)"], 2):
    c = C.cell(row=5, column=i, value=lab); c.font, c.fill, c.border = HDR, HDR_FILL, BOX
rng = f"'Marges produits'!$B${first}:$B${last}"
dv5 = DataValidation(type="list", formula1=f"={rng}", allow_blank=True); C.add_data_validation(dv5)
examples = [("Starter Home", 1), ("Filtre Thermal Vitaminé (Home)", 2)]
for i in range(6):
    r = 6 + i
    nm, q = examples[i] if i < len(examples) else (None, None)
    put(C, f"B{r}", nm, BLUE, fill=YEL); dv5.add(f"B{r}")
    put(C, f"C{r}", q, BLUE, NUM, YEL)
    put(C, f"D{r}", f"=IF(B{r}=\"\",0,C{r}*INDEX('Marges produits'!$E${first}:$E${last},MATCH(B{r},{rng},0)))", GREEN, EUR)
    put(C, f"E{r}", f"=IF(B{r}=\"\",0,C{r}*INDEX('Marges produits'!$G${first}:$G${last},MATCH(B{r},{rng},0)))", GREEN, EUR)
summary = [
    ("Total payé par la cliente, produits (TTC)", "=SUM(D6:D11)"),
    ("Livraison payée par la cliente (TTC)", f"=IF(D13=0,0,{PR['livrfact']})"),
    ("Chiffre d'affaires encaissé (hors TVA si assujettie)", f"=IF({TVA_RECUP},(D13+D14)/(1+{PR['tva']}),D13+D14)"),
    ("− Coût de revient des produits", "=-SUM(E6:E11)"),
    ("− Frais de paiement", f"=IF(D13=0,0,-((D13+D14)*{PR['paypct']}+{PR['payfix']}))"),
    ("− Cotisations sociales", f"=-D15*{PR['cot']}"),
    ("− Expédition", f"=IF(D13=0,0,-{PR['livr']})"),
    ("− Emballage", f"=IF(D13=0,0,-{PR['emb']})"),
    ("− Publicité", f"=IF(D13=0,0,-{PR['pub']})"),
    ("CE QU'IL TE RESTE SUR CETTE COMMANDE", "=SUM(D15:D21)"),
    ("Soit en % du chiffre d'affaires", "=IF(D15=0,0,D22/D15)"),
]
for i, (lab, f) in enumerate(summary):
    r = 13 + i
    put(C, f"B{r}", lab, BOLD if lab.isupper() or lab.startswith("CE ") else BLACK)
    put(C, f"D{r}", f, BOLD if r == 22 else BLACK, PCT if r == 23 else EUR)
C["F13"] = "Astuce : compare un panier d'un seul article avec un panier de 2-3 articles. Les frais fixes (colis, emballage, paiement) pèsent beaucoup moins sur un gros panier."
C["F13"].font, C["F13"].alignment = BLACK, WRAP
C.merge_cells("F13:F17")

# ------------------------------------------------------------------ Seuil de rentabilité
R = sheet("Seuil de rentabilité")
for col, w in zip("ABCD", [3, 46, 16, 60]):
    R.column_dimensions[col].width = w
put(R, "B2", "Frais fixes du mois et seuil de rentabilité", TITLE).border = Border()
R["B3"] = "Mets tes vrais abonnements (exemples en jaune). Mets 0 pour ce que tu n'as pas."
R["B3"].font = BLACK
for i, lab in enumerate(["Frais fixe mensuel", "Montant (€)", "Remarque"], 2):
    c = R.cell(row=5, column=i, value=lab); c.font, c.fill, c.border = HDR, HDR_FILL, BOX
fixed = [
    ("Abonnement Shopify", 36, "Forfait Basic au mois, exemple : vérifie ta facture Shopify."),
    ("Applications Shopify", 0, "Avis, bundles, traduction…"),
    ("Nom de domaine (coût annuel ÷ 12)", 1.5, "Exemple."),
    ("Klaviyo / emailing", 0, "Gratuit jusqu'à un certain nombre de contacts."),
    ("Canva, Metricool, outils de contenu", 0, ""),
    ("Banque professionnelle", 0, ""),
    ("Assurance responsabilité civile pro", 0, "Recommandée pour la vente de produits."),
    ("Comptable / logiciel de facturation", 0, ""),
    ("Stockage, divers", 0, ""),
]
for i, (lab, v, note) in enumerate(fixed):
    r = 6 + i
    put(R, f"B{r}", lab)
    put(R, f"C{r}", v, BLUE, EUR, YEL if v else None)
    put(R, f"D{r}", note)
tot = 6 + len(fixed)
put(R, f"B{tot}", "Total des frais fixes par mois", BOLD)
put(R, f"C{tot}", f"=SUM(C6:C{tot-1})", BOLD, EUR)
put(R, f"B{tot+2}", "Ce qu'il te reste en moyenne par commande (€)", BOLD)
put(R, f"C{tot+2}", "='Simulateur commande'!D22", GREEN, EUR)
put(R, f"D{tot+2}", "Repris du simulateur. Tu peux remplacer par ta vraie moyenne une fois que tu as des ventes.")
put(R, f"B{tot+3}", "Commandes nécessaires par mois pour couvrir les frais fixes", BOLD)
put(R, f"C{tot+3}", f'=IF(C{tot+2}<=0,"Marge négative",ROUNDUP(C{tot}/C{tot+2},0))', BOLD, NUM)
put(R, f"B{tot+4}", "Soit environ par semaine", BLACK)
put(R, f"C{tot+4}", f'=IF(ISNUMBER(C{tot+3}),ROUNDUP(C{tot+3}/4.33,0),"-")', BLACK, NUM)

for wsx in wb.worksheets:
    for row in wsx.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, color=c.font.color, size=c.font.size)

out = __import__("sys").argv[1]
wb.save(out)
print("saved", out)
